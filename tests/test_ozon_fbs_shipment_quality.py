"""Ozon FBS shipment-quality support report helpers."""

from __future__ import annotations

import io
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from review_processor import ozon_fbs_shipment_quality as sq

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (
    Path("/home/ubuntu/.cursor/projects/workspace/uploads/general_fbs_rfbs_rating_8e82.xlsx")
)


def _rating_xlsx(postings: list[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["Отмены (с коэффициентом х2)"])
    ws.append(["Просроченные отгрузки и доставки"])
    ws.append(["*Дата и время указываются Московское UTC +3 часа"])
    ws.append(["Номер отправления"])
    for pn in postings:
        ws.append([pn])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ExtractPostingNumbersTests(unittest.TestCase):
    def test_extracts_after_header_and_skips_dupes(self) -> None:
        raw = _rating_xlsx(
            ["0101152363-0211-1", "0101152363-0211-1", "00214967-1099-7", "not-a-posting"]
        )
        out = sq.extract_posting_numbers_from_rating_xlsx(raw)
        self.assertEqual(out, ["0101152363-0211-1", "00214967-1099-7"])

    def test_requires_header(self) -> None:
        wb = Workbook()
        wb.active.append(["0101152363-0211-1"])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(ValueError):
            sq.extract_posting_numbers_from_rating_xlsx(buf.getvalue())

    def test_sample_upload_file_if_present(self) -> None:
        if not SAMPLE.is_file():
            self.skipTest("sample xlsx not attached")
        out = sq.extract_posting_numbers_from_rating_xlsx(SAMPLE.read_bytes())
        self.assertGreater(len(out), 100)
        self.assertTrue(out[0].count("-") >= 2)


class FormatMoveDateTests(unittest.TestCase):
    def test_msk_date_only(self) -> None:
        # 2026-09-06 21:30 UTC → 07.09.2026 MSK
        dt = datetime(2026, 9, 6, 21, 30, tzinfo=timezone.utc)
        self.assertEqual(sq._format_move_date_only(dt), "07.09.2026")


class BuildSupportRowsTests(unittest.TestCase):
    def test_maps_supply_move_date(self) -> None:
        class _Repo:
            def _sql(self, q: str) -> str:
                return q

            def _row_to_dict(self, r):
                return dict(r)

            def _connect(self):
                repo = self

                class _Cur:
                    def __init__(self, rows):
                        self._rows = rows

                    def fetchall(self):
                        return self._rows

                class _Conn:
                    def __enter__(self):
                        return self

                    def __exit__(self, *a):
                        return False

                    def execute(self, sql, params=()):
                        sql_s = str(sql)
                        if "FROM ozon_fbs_postings" in sql_s:
                            return _Cur(
                                [
                                    {
                                        "posting_number": "A-1-1",
                                        "supply_id": "OZ-1",
                                    },
                                    {
                                        "posting_number": "B-1-1",
                                        "supply_id": "OZ-1",
                                    },
                                ]
                            )
                        if "FROM ozon_fbs_ops_log" in sql_s:
                            return _Cur(
                                [
                                    {
                                        "supply_id": "OZ-1",
                                        "moved_at": datetime(
                                            2026, 9, 6, 10, 0, tzinfo=timezone.utc
                                        ),
                                    }
                                ]
                            )
                        return _Cur([])

                return _Conn()

        with (
            patch.object(sq.oz, "ensure_ozon_fbs_tables"),
            patch.object(sq.ops_log, "ensure_ozon_fbs_ops_log_table"),
        ):
            rows = sq.build_support_report_rows(
                _Repo(),
                user_id=1,
                source_id=18,
                posting_numbers=["A-1-1", "B-1-1", "C-1-1"],
            )
        self.assertEqual(
            rows,
            [
                ("A-1-1", "06.09.2026"),
                ("B-1-1", "06.09.2026"),
                ("C-1-1", ""),
            ],
        )

    def test_xlsx_headers(self) -> None:
        with (
            patch.object(
                sq,
                "build_support_report_rows",
                return_value=[("A-1-1", "06.09.2026")],
            ),
        ):
            payload, fname, meta = sq.build_support_report_xlsx(
                object(),
                user_id=1,
                source_id=18,
                posting_numbers=["A-1-1"],
            )
        self.assertTrue(fname.endswith(".xlsx"))
        self.assertEqual(meta["total"], 1)
        self.assertEqual(meta["with_date"], 1)
        wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
        self.assertEqual(
            list(rows[0]),
            ["Номер отправления", "Дата переноса в доставляются"],
        )
        self.assertEqual(list(rows[1]), ["A-1-1", "06.09.2026"])


class UiReplaceTests(unittest.TestCase):
    def test_quality_button_replaces_sticker_lookup(self) -> None:
        html = (ROOT / "web_templates" / "app.html").read_text(encoding="utf-8")
        js = (ROOT / "web_static" / "ozon_fbs.js").read_text(encoding="utf-8")
        self.assertIn('id="ozonFbsShipmentQualityBtn"', html)
        self.assertIn("Качество отгрузок", html)
        self.assertIn('id="ozonFbsShipmentQualityModal"', html)
        self.assertIn("Загрузить отчет качества отгрузок по ФБС", html)
        self.assertIn("Сформировать отчет для поддержки", html)
        self.assertNotIn('id="ozonFbsStickerLookupBtn"', html)
        self.assertNotIn('id="ozonFbsStickerLookupModal"', html)
        self.assertIn("_ozonFbsSyncOwnerOnlyShipmentQualityBtn", js)
        self.assertIn("/api/ozon-fbs/shipment-quality/support-report", js)
        self.assertIn("ozon_fbs.js?v=137", html)
        self.assertIn("style.css?v=315", html)


if __name__ == "__main__":
    unittest.main()
