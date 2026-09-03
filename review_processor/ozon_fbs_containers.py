"""Ozon FBS carriage containers (грузоместа / паллеты).

Seller API:
- ``POST /v1/carriage/container/create``
- ``POST /v1/carriage/container/list``
- ``POST /v1/carriage/container/get``
- ``POST /v1/carriage/container/fill``
- ``POST /v1/carriage/container/remove-postings``
- ``POST /v1/carriage/container/label/get``
- ``POST /v1/carriage/container/approve``
- ``POST /v1/carriage/container/cancel``  (delete)
"""
from __future__ import annotations

import io
import json
import re
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import UTC, datetime, timedelta
from typing import Any

from . import ozon_fbs as oz
from . import ozon_fbs_supplies as oz_sup
from .repository import ReviewRepository

_log = logging.getLogger(__name__)

SORT_TYPES = frozenset({"sort", "non-sort"})
CARGO_TYPES = frozenset({"box", "pallet"})

# Terminal / already-shipped statuses — hide from the working modal.
_SHIPPED_OR_GONE_STATUSES = frozenset(
    {
        "shipped",
        "closed",
        "cancelled",
        "canceled",
        "deleted",
        "completed",
        "sent",
        "delivered",
    }
)

# Accepted at Ozon sorting center — hidden from the working modal unless explicitly requested.
_SC_ACCEPTED_STATUSES = frozenset({"acceptance_in_progress", "finished"})

_STATUS_LABELS = {
    "new": "Новое",
    "approved": "Подтверждено",
    "approve_failed": "Ошибка подтверждения",
    "acceptance_in_progress": "Принято на СЦ",
    "finished": "Завершено на СЦ",
    "formed": "Сформировано",
    "ready": "Готово",
    "in_progress": "В работе",
    "shipped": "Отгружено",
    "closed": "Закрыто",
    "cancelled": "Отменено",
    "canceled": "Отменено",
    "deleted": "Удалено",
}

_SORT_LABELS = {
    "sort": "Сортируемое",
    "non-sort": "Несортируемое",
}

_CARGO_LABELS = {
    "pallet": "Паллета",
    "box": "Короб",
}


def _utc_now() -> datetime:
    return datetime.now(UTC)


def status_label(status: object) -> str:
    key = str(status or "").strip().lower()
    if not key:
        return "—"
    return _STATUS_LABELS.get(key, key)


# Actions that mean postings can still be added into the cargo place.
_FILL_ACTIONS = frozenset({"fill", "place_posting_into_container"})
# Statuses where Ozon no longer accepts new postings into the container.
_LOCKED_FILL_STATUSES = frozenset(
    {
        "approved",
        "acceptance_in_progress",
        "finished",
        "formed",
        "ready",
        *_SHIPPED_OR_GONE_STATUSES,
    }
)


def is_sc_accepted_container(row: dict[str, Any]) -> bool:
    """True when Ozon has accepted the cargo place at the sorting center."""
    st = str(row.get("status") or "").strip().lower()
    return st in _SC_ACCEPTED_STATUSES


def is_active_container(row: dict[str, Any]) -> bool:
    """True if container should appear in the working modal (not yet shipped)."""
    st = str(row.get("status") or "").strip().lower()
    if st in _SC_ACCEPTED_STATUSES:
        return False
    if st in _SHIPPED_OR_GONE_STATUSES:
        return False
    actions = row.get("available_actions") or []
    if isinstance(actions, list):
        acts = {str(a).strip().lower() for a in actions}
        # Still editable / printable / confirmable → keep.
        if acts & {
            "delete",
            "approve",
            "get_label_container",
            "fill",
            "place_posting_into_container",
        }:
            return True
    # Unknown status without actions: keep unless clearly terminal.
    return st not in _SHIPPED_OR_GONE_STATUSES


def container_accepts_fill(row: dict[str, Any]) -> bool:
    """True when postings can still be scanned/bound into this cargo place."""
    st = str(row.get("status") or "").strip().lower()
    if st in _LOCKED_FILL_STATUSES:
        return False
    actions = row.get("available_actions") or []
    acts_l = {str(a).strip().lower() for a in actions} if isinstance(actions, list) else set()
    if acts_l & _FILL_ACTIONS:
        return True
    # Brand-new / retryable containers are still open even if Ozon omits fill in actions.
    if st in {"new", "approve_failed"} and ("approve" in acts_l or not acts_l):
        return True
    return False


def _normalize_container(raw: dict[str, Any]) -> dict[str, Any] | None:
    if not isinstance(raw, dict):
        return None
    try:
        cid = int(raw.get("container_id") or 0)
    except (TypeError, ValueError):
        cid = 0
    if cid <= 0:
        return None
    try:
        number = int(raw.get("container_number") or 0)
    except (TypeError, ValueError):
        number = 0
    try:
        orders = int(raw.get("count_of_postings") or 0)
    except (TypeError, ValueError):
        orders = 0
    status = str(raw.get("status") or "").strip().lower()
    sort_type = str(raw.get("sort_type") or "").strip().lower()
    cargo_type = str(raw.get("cargo_type") or "").strip().lower()
    actions_raw = raw.get("available_actions") or raw.get("available_actions") or []
    actions = [
        str(a).strip()
        for a in (actions_raw if isinstance(actions_raw, list) else [])
        if str(a).strip()
    ]
    acts_l = {a.lower() for a in actions}
    can_delete = "delete" in acts_l
    # Label is available while Ozon exposes get_label_container (or brand-new «new»).
    can_print = ("get_label_container" in acts_l) or (status in {"new", "approve_failed", "approved"})
    can_approve = "approve" in acts_l and status not in _LOCKED_FILL_STATUSES
    base = {
        "container_id": cid,
        "container_number": number,
        "status": status,
        "status_label": status_label(status),
        "sort_type": sort_type,
        "sort_type_label": _SORT_LABELS.get(sort_type, sort_type or "—"),
        "cargo_type": cargo_type,
        "cargo_type_label": _CARGO_LABELS.get(cargo_type, cargo_type or "—"),
        "order_count": max(0, orders),
        "warehouse_id": raw.get("warehouse_id"),
        "warehouse_name": str(raw.get("warehouse_name") or "").strip(),
        "warehouse_date": str(raw.get("warehouse_date") or "").strip(),
        "created_at": str(raw.get("created_at") or "").strip(),
        "available_actions": actions,
        "can_delete": can_delete,
        "can_print": bool(can_print),
        "can_approve": bool(can_approve),
    }
    base["can_fill"] = container_accepts_fill(base)
    return base


def _friendly_ozon_error(exc: Exception) -> str:
    """Extract Ozon ``message`` from ``Ozon HTTP N: {json}`` RuntimeError text."""
    text = str(exc or "").strip()
    if not text:
        return "Ошибка Ozon API"
    # Typical: Ozon HTTP 400: {"code":3,"message":"FORBIDDEN_TO_CREATE_SORT_BOX"}
    if "{" in text and "}" in text:
        try:
            raw = text[text.index("{") : text.rindex("}") + 1]
            data = json.loads(raw)
            if isinstance(data, dict):
                msg = str(data.get("message") or "").strip()
                if msg:
                    return msg
        except Exception:
            pass
    return text


def resolve_supply_warehouse_id(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
) -> tuple[int, str]:
    """Warehouse for containers: supply row, else first linked posting."""
    sid = str(supply_id or "").strip()
    if not sid:
        raise RuntimeError("Укажите supply_id")
    supply = oz_sup.get_supply(repo, user_id=user_id, source_id=source_id, supply_id=sid)
    wh_name = ""
    wh_id = 0
    if supply:
        wh_name = str(supply.get("warehouse_name") or "").strip()
        try:
            wh = supply.get("warehouse_id")
            wh_id = int(wh) if wh is not None and str(wh).strip() != "" else 0
        except (TypeError, ValueError):
            wh_id = 0
        if wh_id > 0:
            return wh_id, wh_name
    # Fallback: warehouse from any posting linked to this supply.
    oz_sup.ensure_ozon_fbs_supply_schema(repo)
    with repo._connect() as conn:
        row = conn.execute(
            repo._sql(
                """
                SELECT warehouse_id, warehouse_name
                FROM ozon_fbs_postings
                WHERE user_id = ? AND source_id = ? AND supply_id = ?
                  AND warehouse_id IS NOT NULL
                ORDER BY posting_number
                LIMIT 1
                """
            ),
            (user_id, source_id, sid),
        ).fetchone()
    if row:
        d = repo._row_to_dict(row)
        try:
            wh_id = int(d.get("warehouse_id") or 0)
        except (TypeError, ValueError):
            wh_id = 0
        if wh_id > 0:
            return wh_id, str(d.get("warehouse_name") or wh_name).strip() or wh_name
    raise RuntimeError("Не удалось определить склад поставки для грузомест")



def _parse_container_created_at(value: object) -> datetime | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        dt = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return None
    if dt.tzinfo is None:
        return dt.replace(tzinfo=UTC)
    return dt.astimezone(UTC)


def _container_created_sort_key(row: dict[str, Any]) -> tuple:
    """Sort key for newest-first listing (use with ``reverse=True``).

    Parses ``created_at`` to UTC so mixed offsets (``Z`` vs ``+03:00``) order
    correctly. Missing/invalid timestamps sink to the bottom; then higher
    container_number / container_id win.
    """
    created = _parse_container_created_at(row.get("created_at"))
    # Missing/invalid → bottom when sorting reverse=True.
    stamp = created.timestamp() if created is not None else float("-inf")
    try:
        number = int(row.get("container_number") or 0)
    except (TypeError, ValueError):
        number = 0
    try:
        cid = int(row.get("container_id") or 0)
    except (TypeError, ValueError):
        cid = 0
    return (stamp, number, cid)


def list_containers(
    client: oz.OzonFbsClient,
    *,
    warehouse_id: int,
    lookback_days: int = 30,
    include_shipped: bool = False,
    include_sc_accepted: bool = False,
) -> dict[str, Any]:
    wh = int(warehouse_id or 0)
    if wh <= 0:
        raise ValueError("Укажите warehouse_id")
    days = max(1, min(int(lookback_days or 30), 90))
    now = _utc_now()
    created_from = (now - timedelta(days=days)).strftime("%Y-%m-%dT%H:%M:%SZ")
    created_to = now.strftime("%Y-%m-%dT%H:%M:%SZ")
    items: list[dict[str, Any]] = []
    cursor = ""
    for _ in range(20):
        body: dict[str, Any] = {
            "filter": {
                "created_from": created_from,
                "created_to": created_to,
                "warehouse_id": wh,
            },
            "sort_dir": 1,
            "limit": 100,
        }
        if cursor:
            body["cursor"] = cursor
        try:
            data = client.carriage_container_list(body)
        except RuntimeError as exc:
            raise RuntimeError(_friendly_ozon_error(exc)) from exc
        rows = data.get("containers") if isinstance(data, dict) else None
        if not isinstance(rows, list):
            rows = []
        for raw in rows:
            norm = _normalize_container(raw if isinstance(raw, dict) else {})
            if not norm:
                continue
            # Double-check warehouse (API may ignore filter).
            try:
                row_wh = int(norm.get("warehouse_id") or 0)
            except (TypeError, ValueError):
                row_wh = 0
            if row_wh and row_wh != wh:
                continue
            sc_accepted = is_sc_accepted_container(norm)
            if sc_accepted and not include_sc_accepted:
                continue
            if not include_shipped and not is_active_container(norm):
                if not (include_sc_accepted and sc_accepted):
                    continue
            items.append(norm)
        cursor = str((data or {}).get("cursor") or "").strip()
        if not cursor:
            break
    items.sort(key=_container_created_sort_key, reverse=True)
    return {
        "ok": True,
        "warehouse_id": wh,
        "items": items,
        "total": len(items),
        "lookback_days": days,
    }


def create_containers(
    client: oz.OzonFbsClient,
    *,
    warehouse_id: int,
    containers_count: int = 1,
    sort_type: str = "sort",
    cargo_type: str = "pallet",
) -> dict[str, Any]:
    wh = int(warehouse_id or 0)
    if wh <= 0:
        raise ValueError("Укажите warehouse_id")
    try:
        count = int(containers_count)
    except (TypeError, ValueError) as exc:
        raise ValueError("Некорректное число грузомест") from exc
    if count < 1 or count > 100:
        raise ValueError("Можно создать от 1 до 100 грузомест за раз")
    st = str(sort_type or "").strip().lower()
    if st not in SORT_TYPES:
        raise ValueError("sort_type: sort или non-sort")
    ct = str(cargo_type or "").strip().lower()
    if ct not in CARGO_TYPES:
        raise ValueError("cargo_type: box или pallet")
    try:
        data = client.carriage_container_create(
            warehouse_id=wh,
            containers_count=count,
            sort_type=st,
            cargo_type=ct,
        )
    except RuntimeError as exc:
        raise RuntimeError(_friendly_ozon_error(exc)) from exc
    ids_raw = data.get("container_ids") if isinstance(data, dict) else None
    ids: list[int] = []
    if isinstance(ids_raw, list):
        for x in ids_raw:
            try:
                n = int(x)
            except (TypeError, ValueError):
                continue
            if n > 0:
                ids.append(n)
    _log.info(
        "ozon fbs containers create wh=%s count=%s sort=%s cargo=%s ids=%s",
        wh,
        count,
        st,
        ct,
        ids,
    )
    return {
        "ok": True,
        "warehouse_id": wh,
        "container_ids": ids,
        "created": len(ids),
        "sort_type": st,
        "cargo_type": ct,
        "message": f"Создано грузомест: {len(ids)}",
    }


def delete_containers(
    client: oz.OzonFbsClient, *, container_ids: list[int]
) -> dict[str, Any]:
    ids: list[int] = []
    for x in container_ids or []:
        try:
            n = int(x)
        except (TypeError, ValueError):
            continue
        if n > 0:
            ids.append(n)
    ids = ids[:100]
    if not ids:
        raise ValueError("Укажите ID грузомест")
    try:
        data = client.carriage_container_cancel(container_ids=ids)
    except RuntimeError as exc:
        raise RuntimeError(_friendly_ozon_error(exc)) from exc
    errors = []
    if isinstance(data, dict):
        for err in data.get("error_containers") or []:
            if not isinstance(err, dict):
                continue
            errors.append(
                {
                    "container_id": err.get("container_id"),
                    "error": str(err.get("error_message") or err.get("message") or "").strip(),
                }
            )
    ok_n = max(0, len(ids) - len(errors))
    return {
        "ok": not errors,
        "deleted": ok_n,
        "errors": errors,
        "task_id": (data or {}).get("task_id") if isinstance(data, dict) else None,
        "message": (
            f"Удалено: {ok_n}"
            + (f", ошибок {len(errors)}" if errors else "")
        ),
    }


def _wait_container_task(
    client: oz.OzonFbsClient,
    *,
    task_id: int,
    timeout_sec: float = 20.0,
    poll_sec: float = 0.7,
) -> dict[str, Any]:
    """Poll ``/v1/carriage/container/task/info`` until completed/failed or timeout."""
    import time

    tid = int(task_id or 0)
    if tid <= 0:
        return {"ok": True, "status": "", "error_message": ""}
    deadline = time.monotonic() + max(1.0, float(timeout_sec))
    last_status = ""
    last_error = ""
    while time.monotonic() < deadline:
        try:
            info = client.carriage_container_task_info(task_id=tid)
        except Exception as exc:
            _log.warning("ozon container task/info task_id=%s: %s", tid, exc)
            return {
                "ok": False,
                "status": "failed",
                "error_message": _friendly_ozon_error(exc),
            }
        if not isinstance(info, dict):
            info = {}
        last_status = str(info.get("status") or "").strip().lower()
        last_error = str(info.get("error_message") or "").strip()
        if last_status in {"completed", "complete", "success", "done"}:
            return {"ok": True, "status": last_status, "error_message": ""}
        if last_status in {"failed", "error"}:
            return {
                "ok": False,
                "status": last_status,
                "error_message": last_error or "Ошибка подтверждения грузоместа",
            }
        time.sleep(max(0.2, float(poll_sec)))
    # Timeout: not fatal — caller refreshes list; Ozon may still finish.
    return {
        "ok": True,
        "status": last_status or "pending",
        "error_message": last_error,
        "timed_out": True,
    }


def approve_containers(
    client: oz.OzonFbsClient, *, container_ids: list[int]
) -> dict[str, Any]:
    """Confirm cargo-place contents via ``POST /v1/carriage/container/approve``.

    After success Ozon locks the container: no more fill / remove-postings.
    """
    ids: list[int] = []
    for x in container_ids or []:
        try:
            n = int(x)
        except (TypeError, ValueError):
            continue
        if n > 0:
            ids.append(n)
    # API allows a batch; keep a sane upper bound for UI actions.
    ids = ids[:100]
    if not ids:
        raise ValueError("Укажите ID грузомест")
    try:
        data = client.carriage_container_approve(container_ids=ids)
    except RuntimeError as exc:
        raise RuntimeError(_friendly_ozon_error(exc)) from exc
    errors: list[dict[str, Any]] = []
    if isinstance(data, dict):
        for err in data.get("error_containers") or []:
            if not isinstance(err, dict):
                continue
            errors.append(
                {
                    "container_id": err.get("container_id"),
                    "error": str(
                        err.get("error_message") or err.get("message") or ""
                    ).strip(),
                }
            )
    task_id = 0
    if isinstance(data, dict):
        try:
            task_id = int(data.get("task_id") or 0)
        except (TypeError, ValueError):
            task_id = 0
    task_info: dict[str, Any] = {"ok": True, "status": "", "error_message": ""}
    if task_id > 0 and not errors:
        task_info = _wait_container_task(client, task_id=task_id)
        if not task_info.get("ok") and task_info.get("error_message"):
            # Surface async failure as a container-level error when Ozon gave no per-id errors.
            errors.append(
                {
                    "container_id": ids[0] if len(ids) == 1 else None,
                    "error": str(task_info.get("error_message") or "Ошибка подтверждения"),
                }
            )
    ok_n = max(0, len(ids) - len(errors))
    ok = ok_n > 0 and not errors
    if ok and task_info.get("timed_out"):
        message = (
            f"Подтверждение запущено ({ok_n}). Обновите список через несколько секунд."
        )
    elif ok:
        message = f"Подтверждено грузомест: {ok_n}"
    elif ok_n > 0:
        message = f"Подтверждено: {ok_n}, ошибок {len(errors)}"
    else:
        message = (
            "; ".join(e["error"] for e in errors if e.get("error"))
            or "Не удалось подтвердить грузоместо"
        )
    return {
        "ok": ok,
        "approved": ok_n,
        "errors": errors,
        "task_id": task_id or None,
        "task_status": task_info.get("status") or "",
        "message": message,
    }


def build_approve_precheck(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    container_id: int,
) -> dict[str, Any]:
    """Local readiness check before confirming a cargo place.

    - ``has_sync_errors``: postings bound to this container with Ozon sync errors
      (requires explicit force to approve).
    - ``has_unbound``: supply still has postings without any cargo place
      (warning only).
    """
    try:
        cid = int(container_id)
    except (TypeError, ValueError) as exc:
        raise ValueError("Некорректный container_id") from exc
    if cid <= 0:
        raise ValueError("Укажите ID грузоместа")
    supply = oz_sup.get_supply(
        repo, user_id=user_id, source_id=source_id, supply_id=str(supply_id)
    )
    nums = [
        str(x).strip()
        for x in ((supply or {}).get("posting_numbers") or [])
        if str(x).strip()
    ]
    binds = load_container_bind_map(
        repo, user_id=user_id, source_id=source_id, posting_numbers=nums
    )
    sync_errors: list[dict[str, str]] = []
    bound_here = 0
    bound_other = 0
    unbound = 0
    for pn in nums:
        row = binds.get(pn) or {}
        try:
            row_cid = int(row.get("container_id") or 0)
        except (TypeError, ValueError):
            row_cid = 0
        err = str(row.get("container_sync_error") or "").strip()
        if row_cid == cid:
            bound_here += 1
            if err:
                sync_errors.append({"posting_number": pn, "error": err})
        elif row_cid > 0:
            bound_other += 1
        else:
            unbound += 1
    has_sync_errors = bool(sync_errors)
    has_unbound = unbound > 0
    return {
        "ok": True,
        "container_id": cid,
        "total_orders": len(nums),
        "bound_to_container": bound_here,
        "bound_other": bound_other,
        "unbound": unbound,
        "sync_error_count": len(sync_errors),
        "sync_errors": sync_errors[:20],
        "has_sync_errors": has_sync_errors,
        "has_unbound": has_unbound,
        "requires_force": has_sync_errors,
        "message": (
            (
                f"Ошибки синхронизации с Ozon: {len(sync_errors)}. "
                "Состав на портале может отличаться."
            )
            if has_sync_errors
            else (
                f"Не привязано к грузоместам: {unbound} из {len(nums)}."
                if has_unbound
                else ""
            )
        ),
    }


def get_container_labels_pdf(
    client: oz.OzonFbsClient, *, container_ids: list[int]
) -> dict[str, Any]:
    ids: list[int] = []
    for x in container_ids or []:
        try:
            n = int(x)
        except (TypeError, ValueError):
            continue
        if n > 0:
            ids.append(n)
    ids = ids[:300]
    if not ids:
        raise ValueError("Укажите ID грузомест")
    try:
        data = client.carriage_container_label_get(container_ids=ids)
    except RuntimeError as exc:
        raise RuntimeError(_friendly_ozon_error(exc)) from exc
    content = data.get("content") if isinstance(data, dict) else None
    file_b64 = ""
    content_type = "application/pdf"
    if isinstance(content, dict):
        file_b64 = str(content.get("file_content") or "").strip()
        content_type = str(content.get("content_type") or content_type).strip() or content_type
    errors = []
    if isinstance(data, dict):
        for err in data.get("error_containers") or []:
            if not isinstance(err, dict):
                continue
            errors.append(
                {
                    "container_id": err.get("container_id"),
                    "error": str(err.get("error_message") or "").strip(),
                }
            )
    if not file_b64:
        detail = "; ".join(e["error"] for e in errors if e.get("error")) or "Ozon не вернул PDF"
        raise RuntimeError(detail)
    return {
        "ok": True,
        "content_type": content_type,
        "file_content": file_b64,
        "errors": errors,
        "container_ids": ids,
    }


def normalize_container_scan(value: object) -> str:
    """Normalize scanned cargo QR (digits from printed label)."""
    raw = str(value or "").strip()
    if not raw:
        return ""
    digits = re.sub(r"\D+", "", raw)
    return digits or raw


def match_container_by_scan(
    containers: list[dict[str, Any]], scan: object
) -> dict[str, Any] | None:
    """Match scanned QR to a known container (by container_id string primarily)."""
    key = normalize_container_scan(scan)
    if not key:
        return None
    for row in containers or []:
        if not isinstance(row, dict):
            continue
        cid = str(row.get("container_id") or "").strip()
        barcode = normalize_container_scan(row.get("container_barcode") or row.get("barcode") or "")
        number = str(row.get("container_number") or "").strip()
        if cid and cid == key:
            return row
        if barcode and barcode == key:
            return row
        if number and number == key and len(key) >= 6:
            return row
    return None


def load_container_bind_map(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_numbers: list[str],
) -> dict[str, dict[str, Any]]:
    nums = [str(x).strip() for x in posting_numbers if str(x).strip()]
    if not nums:
        return {}
    oz.ensure_ozon_fbs_tables(repo)
    placeholders = ", ".join("?" for _ in nums)
    with repo._connect() as conn:
        rows = conn.execute(
            repo._sql(
                f"""
                SELECT posting_number, container_id, container_barcode,
                       container_synced, container_sync_error
                FROM ozon_fbs_postings
                WHERE user_id = ? AND source_id = ?
                  AND posting_number IN ({placeholders})
                """
            ),
            (user_id, source_id, *nums),
        ).fetchall()
    out: dict[str, dict[str, Any]] = {}
    for row in rows:
        d = repo._row_to_dict(row)
        pn = str(d.get("posting_number") or "").strip()
        if not pn:
            continue
        try:
            cid = int(d.get("container_id") or 0)
        except (TypeError, ValueError):
            cid = 0
        barcode = str(d.get("container_barcode") or "").strip()
        if not barcode and cid > 0:
            barcode = str(cid)
        out[pn] = {
            "container_id": cid if cid > 0 else None,
            "container_barcode": barcode,
            "container_synced": bool(d.get("container_synced")) and cid > 0,
            "container_sync_error": str(d.get("container_sync_error") or "").strip(),
        }
    return out


def _set_local_container_bind(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_number: str,
    container_id: int | None,
    container_barcode: str = "",
    synced: bool = False,
    sync_error: str = "",
) -> dict[str, Any]:
    oz.ensure_ozon_fbs_tables(repo)
    pn = str(posting_number or "").strip()
    if not pn:
        raise ValueError("Не указан номер отправления")
    cid = int(container_id or 0)
    barcode = normalize_container_scan(container_barcode) or (str(cid) if cid > 0 else "")
    with repo._connect() as conn:
        conn.execute(
            repo._sql(
                """
                UPDATE ozon_fbs_postings
                SET container_id = ?,
                    container_barcode = ?,
                    container_synced = ?,
                    container_sync_error = ?
                WHERE user_id = ? AND source_id = ? AND posting_number = ?
                """
            ),
            (
                cid if cid > 0 else None,
                barcode if cid > 0 else "",
                bool(synced) and cid > 0,
                str(sync_error or "").strip()[:500],
                user_id,
                source_id,
                pn,
            ),
        )
    return {
        "posting_number": pn,
        "container_id": cid if cid > 0 else None,
        "container_barcode": barcode if cid > 0 else "",
        "container_synced": bool(synced) and cid > 0,
        "container_sync_error": str(sync_error or "").strip(),
    }


def _fetch_container_row(
    client: oz.OzonFbsClient, *, container_id: int
) -> dict[str, Any] | None:
    """Normalize one container from ``/v1/carriage/container/get`` (best-effort)."""
    try:
        data = client.carriage_container_get(container_id=int(container_id))
    except Exception as exc:
        _log.warning("ozon container get cid=%s: %s", container_id, exc)
        return None
    raw = data
    if isinstance(data, dict):
        nested = data.get("container") or data.get("result")
        if isinstance(nested, dict):
            raw = nested
    if not isinstance(raw, dict):
        return None
    # Ensure id present for normalize when Ozon omits it in nested payload.
    if not raw.get("container_id"):
        raw = {**raw, "container_id": int(container_id)}
    return _normalize_container(raw)


def bind_posting_to_container(
    client: oz.OzonFbsClient,
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_number: str,
    container_id: int,
    container_barcode: str = "",
    previous_container_id: int | None = None,
) -> dict[str, Any]:
    """Local bind + Ozon fill (and remove from previous if needed).

    Local binding is kept even when Ozon fails; sync_error is stored.
    Confirmed (approved) cargo places reject new fills before calling Ozon.
    """
    pn = str(posting_number or "").strip()
    try:
        cid = int(container_id)
    except (TypeError, ValueError) as exc:
        raise ValueError("Некорректный container_id") from exc
    if not pn or cid <= 0:
        raise ValueError("Укажите отправление и грузоместо")
    barcode = normalize_container_scan(container_barcode) or str(cid)
    # Guard: confirmed / locked containers must not accept new postings.
    meta = _fetch_container_row(client, container_id=cid)
    if meta is not None and not container_accepts_fill(meta):
        raise ValueError(
            f"Грузоместо {cid} уже подтверждено — в него нельзя добавить заказы"
        )
    prev = int(previous_container_id or 0)
    sync_error = ""
    synced = False
    try:
        if prev > 0 and prev != cid:
            try:
                client.carriage_container_remove_postings(
                    container_id=prev, posting_numbers=[pn]
                )
            except Exception as rem_exc:
                _log.warning(
                    "ozon container remove-postings prev=%s pn=%s: %s",
                    prev,
                    pn,
                    rem_exc,
                )
        client.carriage_container_fill(container_id=cid, posting_numbers=[pn])
        synced = True
    except Exception as exc:
        sync_error = _friendly_ozon_error(exc)
        _log.warning("ozon container fill cid=%s pn=%s: %s", cid, pn, sync_error)
    local = _set_local_container_bind(
        repo,
        user_id=user_id,
        source_id=source_id,
        posting_number=pn,
        container_id=cid,
        container_barcode=barcode,
        synced=synced,
        sync_error=sync_error,
    )
    return {
        "ok": True,
        "synced": synced,
        "error": sync_error,
        **local,
    }


def unbind_posting_from_container(
    client: oz.OzonFbsClient | None,
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_number: str,
    container_id: int | None = None,
) -> dict[str, Any]:
    """Clear local bind; optionally remove posting from Ozon container."""
    pn = str(posting_number or "").strip()
    if not pn:
        raise ValueError("Не указан номер отправления")
    existing = load_container_bind_map(
        repo, user_id=user_id, source_id=source_id, posting_numbers=[pn]
    ).get(pn) or {}
    cid = int(container_id or existing.get("container_id") or 0)
    if client is not None and cid > 0:
        try:
            client.carriage_container_remove_postings(
                container_id=cid, posting_numbers=[pn]
            )
        except Exception as exc:
            # Local clear still proceeds; Ozon mismatch is only logged.
            _log.warning(
                "ozon container remove-postings cid=%s pn=%s: %s",
                cid,
                pn,
                _friendly_ozon_error(exc),
            )
    local = _set_local_container_bind(
        repo,
        user_id=user_id,
        source_id=source_id,
        posting_number=pn,
        container_id=None,
        container_barcode="",
        synced=False,
        sync_error="",
    )
    return {"ok": True, "error": "", **local}


_CONTAINER_LIST_CACHE: dict[tuple[int, int, int], tuple[float, dict[str, Any]]] = {}
_CONTAINER_LIST_CACHE_TTL_SEC = 45.0
_RECONCILE_FETCH_WORKERS = 8


def _posting_numbers_from_container_payload(raw: dict[str, Any]) -> list[str]:
    """Extract posting numbers from ``/v1/carriage/container/get`` payload (best-effort)."""
    out: list[str] = []
    seen: set[str] = set()

    def add(value: object) -> None:
        pn = str(value or "").strip()
        if pn and pn not in seen:
            seen.add(pn)
            out.append(pn)

    direct = raw.get("posting_numbers")
    if isinstance(direct, list):
        for item in direct:
            if isinstance(item, str):
                add(item)
            elif isinstance(item, dict):
                add(item.get("posting_number"))
    postings = raw.get("postings")
    if isinstance(postings, list):
        for item in postings:
            if isinstance(item, str):
                add(item)
            elif isinstance(item, dict):
                add(item.get("posting_number"))
    return out


def _container_get_raw_payload(data: object) -> dict[str, Any]:
    if not isinstance(data, dict):
        return {}
    nested = data.get("container") or data.get("result")
    if isinstance(nested, dict):
        return nested
    return data


def _fetch_container_postings(
    client: oz.OzonFbsClient, *, container_id: int
) -> tuple[dict[str, Any] | None, list[str], bool]:
    """Return (normalized meta, posting_numbers, exists_on_ozon)."""
    try:
        data = client.carriage_container_get(container_id=int(container_id))
    except Exception as exc:
        text = str(exc or "").casefold()
        if "404" in text or "not found" in text or "not_found" in text:
            return None, [], False
        _log.warning("ozon container get cid=%s: %s", container_id, exc)
        return None, [], True
    raw = _container_get_raw_payload(data)
    if not raw:
        return None, [], False
    if not raw.get("container_id"):
        raw = {**raw, "container_id": int(container_id)}
    meta = _normalize_container(raw)
    postings = _posting_numbers_from_container_payload(raw)
    return meta, postings, True


def _list_containers_cached(
    client: oz.OzonFbsClient,
    *,
    user_id: int,
    source_id: int,
    warehouse_id: int,
    lookback_days: int = 30,
) -> dict[str, Any]:
    key = (int(user_id), int(source_id), int(warehouse_id))
    now = _utc_now().timestamp()
    cached = _CONTAINER_LIST_CACHE.get(key)
    if cached and (now - cached[0]) < _CONTAINER_LIST_CACHE_TTL_SEC:
        return cached[1]
    listed = list_containers(
        client, warehouse_id=int(warehouse_id), lookback_days=lookback_days
    )
    _CONTAINER_LIST_CACHE[key] = (now, listed)
    return listed


def reconcile_supply_container_binds(
    client: oz.OzonFbsClient,
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    skip_postings: list[str] | None = None,
) -> dict[str, Any]:
    """Align local posting→container binds with Ozon portal (portal wins).

    Read-only towards Ozon except local DB updates — no fill/remove during reconcile.
    """
    sid = str(supply_id or "").strip()
    if not sid:
        raise ValueError("Укажите supply_id")
    skip = {str(x).strip() for x in (skip_postings or []) if str(x).strip()}

    supply = oz_sup.get_supply(
        repo, user_id=user_id, source_id=source_id, supply_id=sid
    )
    posting_numbers = [
        str(x).strip()
        for x in ((supply or {}).get("posting_numbers") or [])
        if str(x).strip()
    ]
    supply_set = set(posting_numbers)
    if not supply_set:
        return {
            "ok": True,
            "changes": [],
            "binds": {},
            "posting_lists_available": False,
            "container_reconciled_at": _utc_now().isoformat(),
            "message": "В поставке нет отправлений",
        }

    local_binds = load_container_bind_map(
        repo, user_id=user_id, source_id=source_id, posting_numbers=posting_numbers
    )

    warehouse_id, _wh_name = resolve_supply_warehouse_id(
        repo, user_id=user_id, source_id=source_id, supply_id=sid
    )
    listed = _list_containers_cached(
        client,
        user_id=user_id,
        source_id=source_id,
        warehouse_id=warehouse_id,
    )
    alive_ids: set[int] = set()
    alive_meta: dict[int, dict[str, Any]] = {}
    for item in listed.get("items") or []:
        if not isinstance(item, dict):
            continue
        try:
            cid = int(item.get("container_id") or 0)
        except (TypeError, ValueError):
            cid = 0
        if cid > 0:
            alive_ids.add(cid)
            alive_meta[cid] = item

    local_cids: set[int] = set()
    for row in local_binds.values():
        try:
            cid = int(row.get("container_id") or 0)
        except (TypeError, ValueError):
            cid = 0
        if cid > 0:
            local_cids.add(cid)

    fetch_ids = set(alive_ids) | local_cids
    ozon_map: dict[str, int] = {}
    posting_lists_available = False
    missing_container_ids: set[int] = set()

    if fetch_ids:
        workers = min(_RECONCILE_FETCH_WORKERS, len(fetch_ids))
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            futures = {
                pool.submit(_fetch_container_postings, client, container_id=cid): cid
                for cid in fetch_ids
            }
            for fut in as_completed(futures):
                cid = futures[fut]
                try:
                    meta, nums, exists = fut.result()
                except Exception as exc:
                    _log.warning("reconcile container get cid=%s: %s", cid, exc)
                    continue
                if not exists:
                    missing_container_ids.add(int(cid))
                    continue
                if meta is not None and int(cid) not in alive_meta:
                    alive_meta[int(cid)] = meta
                if nums:
                    posting_lists_available = True
                    for pn in nums:
                        if pn in supply_set:
                            ozon_map[pn] = int(cid)

    changes: list[dict[str, Any]] = []

    def _apply_local(
        pn: str,
        *,
        container_id: int | None,
        synced: bool,
        sync_error: str,
        action: str,
        reason: str,
    ) -> None:
        barcode = str(container_id) if container_id else ""
        _set_local_container_bind(
            repo,
            user_id=user_id,
            source_id=source_id,
            posting_number=pn,
            container_id=container_id,
            container_barcode=barcode,
            synced=synced,
            sync_error=sync_error,
        )
        changes.append(
            {
                "posting_number": pn,
                "action": action,
                "reason": reason,
                "container_id": container_id,
                "container_barcode": barcode if container_id else "",
                "container_synced": bool(synced) and bool(container_id),
                "container_sync_error": sync_error,
            }
        )

    for pn in posting_numbers:
        if pn in skip:
            continue
        local = local_binds.get(pn) or {}
        try:
            local_cid = int(local.get("container_id") or 0)
        except (TypeError, ValueError):
            local_cid = 0
        local_synced = bool(local.get("container_synced")) and local_cid > 0
        local_err = str(local.get("container_sync_error") or "").strip()
        ozon_cid = ozon_map.get(pn)

        if local_cid > 0 and local_cid in missing_container_ids:
            if not local_synced and local_err:
                continue
            _apply_local(
                pn,
                container_id=None,
                synced=False,
                sync_error="",
                action="cleared",
                reason="Грузоместо удалено на портале Ozon",
            )
            continue

        if not posting_lists_available:
            continue

        if ozon_cid == local_cid:
            if ozon_cid > 0 and not local_synced and not local_err:
                _apply_local(
                    pn,
                    container_id=ozon_cid,
                    synced=True,
                    sync_error="",
                    action="synced",
                    reason="Привязка подтверждена на портале Ozon",
                )
            continue

        if ozon_cid is None or ozon_cid <= 0:
            if local_cid <= 0:
                continue
            if not local_synced and local_err:
                continue
            _apply_local(
                pn,
                container_id=None,
                synced=False,
                sync_error="",
                action="cleared",
                reason="На портале Ozon заказ не в грузоместе",
            )
            continue

        action = "adopted" if local_cid <= 0 else "updated"
        _apply_local(
            pn,
            container_id=ozon_cid,
            synced=True,
            sync_error="",
            action=action,
            reason="Состав грузоместа синхронизирован с порталом Ozon",
        )

    updated_binds = load_container_bind_map(
        repo, user_id=user_id, source_id=source_id, posting_numbers=posting_numbers
    )
    return {
        "ok": True,
        "changes": changes,
        "binds": updated_binds,
        "posting_lists_available": posting_lists_available,
        "containers_checked": len(fetch_ids),
        "container_reconciled_at": _utc_now().isoformat(),
        "message": (
            f"Синхронизировано: {len(changes)}"
            if changes
            else (
                "Данные совпадают с порталом"
                if posting_lists_available
                else "Состав грузомест недоступен из API — проверены только удалённые ГМ"
            )
        ),
    }


def container_bind_fields_from_map(
    bind_map: dict[str, dict[str, Any]], posting_number: str
) -> dict[str, Any]:
    row = bind_map.get(str(posting_number or "").strip()) or {}
    cid = row.get("container_id")
    try:
        cid_i = int(cid or 0)
    except (TypeError, ValueError):
        cid_i = 0
    barcode = str(row.get("container_barcode") or "").strip()
    if not barcode and cid_i > 0:
        barcode = str(cid_i)
    err = str(row.get("container_sync_error") or "").strip()
    return {
        "container_id": cid_i if cid_i > 0 else None,
        "container_barcode": barcode,
        "container_synced": bool(row.get("container_synced")) and cid_i > 0 and not err,
        "container_sync_error": err,
    }


def get_supply_moved_to_delivering_at(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
) -> str:
    """Latest local «move to Доставляются» timestamp for the supply (ISO), or ``""``."""
    from . import ozon_fbs_ops_log as ops_log

    sid = str(supply_id or "").strip()
    if not sid:
        return ""
    try:
        ops_log.ensure_ozon_fbs_ops_log_table(repo)
        with repo._connect() as conn:
            row = conn.execute(
                repo._sql(
                    """
                    SELECT created_at
                    FROM ozon_fbs_ops_log
                    WHERE user_id = ?
                      AND source_id = ?
                      AND supply_id = ?
                      AND action = ?
                    ORDER BY created_at DESC, id DESC
                    LIMIT 1
                    """
                ),
                (
                    int(user_id),
                    int(source_id),
                    sid,
                    ops_log.ACTION_MOVE_DELIVERING,
                ),
            ).fetchone()
    except Exception as exc:
        _log.warning(
            "ozon fbs move-delivering lookup failed supply=%s: %s", sid, exc
        )
        return ""
    if not row:
        return ""
    d = repo._row_to_dict(row)
    created = d.get("created_at")
    if created is None:
        return ""
    if isinstance(created, datetime):
        if created.tzinfo is None:
            created = created.replace(tzinfo=UTC)
        return created.astimezone(UTC).isoformat()
    return str(created).strip()


def enrich_containers_for_supply_modal(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    listed: dict[str, Any],
) -> dict[str, Any]:
    """Cheap list enrichment: display dates + one supply-level delivering move time.

    Does not call Ozon beyond the already-fetched ``list`` payload.
    """
    moved_raw = get_supply_moved_to_delivering_at(
        repo,
        user_id=user_id,
        source_id=source_id,
        supply_id=supply_id,
    )
    moved_display = oz.format_lookup_datetime(moved_raw) if moved_raw else ""
    items_in = listed.get("items") if isinstance(listed, dict) else None
    items: list[dict[str, Any]] = []
    if isinstance(items_in, list):
        for raw in items_in:
            if not isinstance(raw, dict):
                continue
            row = dict(raw)
            wh = row.get("warehouse_date")
            created = row.get("created_at")
            row["warehouse_date_display"] = oz.format_warehouse_date(wh) if wh else ""
            row["created_at_display"] = (
                oz.format_lookup_datetime(created) if created else ""
            )
            row["moved_to_delivering_at"] = moved_raw
            row["moved_to_delivering_at_display"] = moved_display
            items.append(row)
    out = dict(listed) if isinstance(listed, dict) else {"ok": True, "items": []}
    out["items"] = items
    out["moved_to_delivering_at"] = moved_raw
    out["moved_to_delivering_at_display"] = moved_display
    return out


def _list_local_container_postings(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    container_id: int,
    supply_id: str = "",
) -> list[dict[str, Any]]:
    """Postings bound locally to this cargo place (KIZ / pick-verify binds)."""
    oz.ensure_ozon_fbs_tables(repo)
    cid = int(container_id or 0)
    if cid <= 0:
        return []
    sid = str(supply_id or "").strip()
    clauses = [
        "user_id = ?",
        "source_id = ?",
        "container_id = ?",
    ]
    params: list[Any] = [int(user_id), int(source_id), cid]
    if sid:
        clauses.append("supply_id = ?")
        params.append(sid)
    where = " AND ".join(clauses)
    with repo._connect() as conn:
        rows = conn.execute(
            repo._sql(
                f"""
                SELECT posting_number, offer_id, product_name, quantity,
                       marking_codes_json, pick_verified, pick_barcode,
                       container_barcode, container_synced, tab, status
                FROM ozon_fbs_postings
                WHERE {where}
                ORDER BY posting_number
                """
            ),
            tuple(params),
        ).fetchall()
    out = _rows_to_container_posting_items(repo, rows)
    # Binds may exist with empty/stale supply_id after moves — don't hide them.
    if not out and sid:
        return _list_local_container_postings(
            repo,
            user_id=user_id,
            source_id=source_id,
            container_id=cid,
            supply_id="",
        )
    return out


def _rows_to_container_posting_items(
    repo: ReviewRepository, rows: list[Any]
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in rows:
        d = repo._row_to_dict(row)
        pn = str(d.get("posting_number") or "").strip()
        if not pn:
            continue
        try:
            codes = json.loads(str(d.get("marking_codes_json") or "[]"))
        except (TypeError, ValueError, json.JSONDecodeError):
            codes = []
        if not isinstance(codes, list):
            codes = []
        kiz_n = sum(1 for x in codes if str(x or "").strip())
        pick_barcode = str(d.get("pick_barcode") or "").strip()
        pick_ok = bool(d.get("pick_verified")) and bool(pick_barcode)
        try:
            qty = int(d.get("quantity") or 1)
        except (TypeError, ValueError):
            qty = 1
        out.append(
            {
                "posting_number": pn,
                "offer_id": str(d.get("offer_id") or "").strip(),
                "product_name": str(d.get("product_name") or "").strip(),
                "quantity": max(1, qty),
                "has_kiz": kiz_n > 0,
                "kiz_count": kiz_n,
                "pick_verified": pick_ok,
                "container_barcode": str(d.get("container_barcode") or "").strip(),
                "container_synced": bool(d.get("container_synced")),
                "tab": str(d.get("tab") or "").strip(),
                "status": str(d.get("status") or "").strip(),
            }
        )
    return out


def _load_posting_items_by_numbers(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_numbers: list[str],
) -> dict[str, dict[str, Any]]:
    """Map posting_number → composition row fields from local DB."""
    nums = [str(x).strip() for x in posting_numbers if str(x).strip()]
    if not nums:
        return {}
    oz.ensure_ozon_fbs_tables(repo)
    out: dict[str, dict[str, Any]] = {}
    chunk = 400
    with repo._connect() as conn:
        for i in range(0, len(nums), chunk):
            part = nums[i : i + chunk]
            placeholders = ", ".join("?" for _ in part)
            rows = conn.execute(
                repo._sql(
                    f"""
                    SELECT posting_number, offer_id, product_name, quantity,
                           marking_codes_json, pick_verified, pick_barcode,
                           container_barcode, container_synced, tab, status
                    FROM ozon_fbs_postings
                    WHERE user_id = ? AND source_id = ?
                      AND posting_number IN ({placeholders})
                    """
                ),
                (int(user_id), int(source_id), *part),
            ).fetchall()
            for item in _rows_to_container_posting_items(repo, rows):
                out[str(item["posting_number"])] = item
    return out


def _stub_container_posting_item(posting_number: str) -> dict[str, Any]:
    pn = str(posting_number or "").strip()
    return {
        "posting_number": pn,
        "offer_id": "",
        "product_name": "",
        "quantity": 1,
        "has_kiz": False,
        "kiz_count": 0,
        "pick_verified": False,
        "container_barcode": "",
        "container_synced": False,
        "tab": "",
        "status": "",
    }


def _merge_ozon_container_composition(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    container_id: int,
    ozon_posting_numbers: list[str],
    local_postings: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Prefer Ozon container membership; enrich rows from local posting fields."""
    ozon_pns = [str(x).strip() for x in ozon_posting_numbers if str(x).strip()]
    if not ozon_pns:
        return list(local_postings or [])
    by_local = {
        str(p.get("posting_number") or "").strip(): p
        for p in (local_postings or [])
        if str(p.get("posting_number") or "").strip()
    }
    loaded = _load_posting_items_by_numbers(
        repo,
        user_id=user_id,
        source_id=source_id,
        posting_numbers=ozon_pns,
    )
    supply = oz_sup.get_supply(
        repo, user_id=user_id, source_id=source_id, supply_id=str(supply_id or "")
    )
    supply_set = {
        str(x).strip()
        for x in ((supply or {}).get("posting_numbers") or [])
        if str(x).strip()
    }
    cid = int(container_id or 0)
    # Soft-fill missing local binds for supply postings found on Ozon (portal wins
    # only when local has no container yet — do not fight an active edit).
    if cid > 0 and supply_set:
        candidates = [pn for pn in ozon_pns if pn in supply_set]
        binds = load_container_bind_map(
            repo,
            user_id=user_id,
            source_id=source_id,
            posting_numbers=candidates,
        )
        for pn in candidates:
            bind = binds.get(pn) or {}
            try:
                local_cid = int(bind.get("container_id") or 0)
            except (TypeError, ValueError):
                local_cid = 0
            if local_cid > 0:
                continue
            try:
                _set_local_container_bind(
                    repo,
                    user_id=user_id,
                    source_id=source_id,
                    posting_number=pn,
                    container_id=cid,
                    container_barcode=str(cid),
                    synced=True,
                    sync_error="",
                )
            except Exception as exc:
                _log.warning(
                    "ozon container compose soft-bind %s→%s: %s", pn, cid, exc
                )

    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    for pn in ozon_pns:
        if pn in seen:
            continue
        seen.add(pn)
        row = dict(loaded.get(pn) or by_local.get(pn) or _stub_container_posting_item(pn))
        row["posting_number"] = pn
        if supply_set and pn not in supply_set:
            row["outside_supply"] = True
        out.append(row)
    return out


def build_container_modal_details(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    container: dict[str, Any],
    client: oz.OzonFbsClient | None = None,
) -> dict[str, Any]:
    """Timeline + composition for one cargo place.

    Local binds first; when ``client`` is set, one ``container/get`` fills the
    real Ozon membership (fixes «Заказов 195 / Состав 0» when binds were not
    stored locally).
    """
    try:
        cid = int(container.get("container_id") or 0)
    except (TypeError, ValueError):
        cid = 0
    status = str(container.get("status") or "").strip().lower()
    status_lbl = str(container.get("status_label") or status_label(status) or "—")
    created_raw = str(container.get("created_at") or "").strip()
    warehouse_raw = str(container.get("warehouse_date") or "").strip()

    local_postings = _list_local_container_postings(
        repo,
        user_id=user_id,
        source_id=source_id,
        container_id=cid,
        supply_id=supply_id,
    )
    composition_source = "local"
    ozon_fetch_ok = False
    postings = list(local_postings)
    if client is not None and cid > 0:
        try:
            meta_get, ozon_pns, exists = _fetch_container_postings(
                client, container_id=cid
            )
            ozon_fetch_ok = bool(exists)
            # ``container/get`` may carry a richer warehouse_date / created_at than list.
            if isinstance(meta_get, dict):
                wh_get = str(meta_get.get("warehouse_date") or "").strip()
                if wh_get and (len(wh_get) > len(warehouse_raw) or not warehouse_raw):
                    warehouse_raw = wh_get
                cr_get = str(meta_get.get("created_at") or "").strip()
                if cr_get and (len(cr_get) > len(created_raw) or not created_raw):
                    created_raw = cr_get
                st_get = str(meta_get.get("status") or "").strip().lower()
                if st_get:
                    status = st_get
                    status_lbl = str(
                        meta_get.get("status_label") or status_label(status) or "—"
                    )
            if ozon_pns:
                postings = _merge_ozon_container_composition(
                    repo,
                    user_id=user_id,
                    source_id=source_id,
                    supply_id=supply_id,
                    container_id=cid,
                    ozon_posting_numbers=ozon_pns,
                    local_postings=local_postings,
                )
                composition_source = "ozon"
            elif exists and not local_postings:
                composition_source = "ozon"
        except Exception as exc:
            _log.warning("ozon container details get cid=%s: %s", cid, exc)
            ozon_fetch_ok = False

    created_display = oz.format_lookup_datetime(created_raw) if created_raw else ""
    warehouse_display = (
        oz.format_warehouse_date(warehouse_raw) if warehouse_raw else ""
    )
    moved_raw = get_supply_moved_to_delivering_at(
        repo,
        user_id=user_id,
        source_id=source_id,
        supply_id=supply_id,
    )
    moved_display = oz.format_lookup_datetime(moved_raw) if moved_raw else ""

    timeline: list[dict[str, Any]] = []
    if created_raw or created_display:
        timeline.append(
            {
                "key": "created",
                "label": "Создано",
                "at": created_raw,
                "at_display": created_display or "—",
                "source": "ozon",
            }
        )
    # Local ops_log: when the supply (with this GM) was moved to «Доставляются».
    timeline.append(
        {
            "key": "moved_to_delivering",
            "label": "Дата отгрузки с нашего склада",
            "at": moved_raw,
            "at_display": moved_display or "—",
            "source": "local",
        }
    )
    # Docs: warehouse_date = creation date in warehouse TZ (string). Ozon often
    # sends YYYY-MM-DD only; we show HH:MM whenever the payload includes time.
    if warehouse_raw or warehouse_display:
        timeline.append(
            {
                "key": "warehouse_date",
                "label": "Дата склада (Ozon)",
                "at": warehouse_raw,
                "at_display": warehouse_display or "—",
                "source": "ozon",
            }
        )
    timeline.append(
        {
            "key": "status",
            "label": f"Текущий статус: {status_lbl}",
            "at": "",
            "at_display": status_lbl,
            "source": "ozon",
            "status": status,
            "status_label": status_lbl,
        }
    )

    return {
        "ok": True,
        "container_id": cid,
        "container_number": container.get("container_number"),
        "status": status,
        "status_label": status_lbl,
        "warehouse_date": warehouse_raw,
        "warehouse_date_display": warehouse_display,
        "created_at": created_raw,
        "created_at_display": created_display,
        "moved_to_delivering_at": moved_raw,
        "moved_to_delivering_at_display": moved_display,
        "timeline": timeline,
        "postings": postings,
        "postings_count": len(postings),
        "composition_source": composition_source,
        "ozon_fetch_ok": ozon_fetch_ok,
        "local_postings_count": len(local_postings),
    }


def get_container_modal_details(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    container_id: int,
    container: dict[str, Any] | None = None,
    client: oz.OzonFbsClient | None = None,
) -> dict[str, Any]:
    """Build GM expand payload: timeline from list-row meta + composition.

    When ``client`` is provided, one Ozon ``container/get`` loads membership so
    «Состав» matches «Заказов». Without client — local binds only (tests / offline).
    """
    cid = int(container_id or 0)
    if cid <= 0:
        raise ValueError("Укажите container_id")
    sid = str(supply_id or "").strip()
    if not sid:
        raise ValueError("Укажите supply_id")
    meta = dict(container) if isinstance(container, dict) else {}
    if not meta.get("container_id"):
        meta["container_id"] = cid
    return build_container_modal_details(
        repo,
        user_id=user_id,
        source_id=source_id,
        supply_id=sid,
        container=meta,
        client=client,
    )


def build_container_composition_xlsx(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_id: str,
    container_id: int,
    container: dict[str, Any] | None = None,
    client: oz.OzonFbsClient | None = None,
) -> tuple[bytes, str]:
    """XLSX: posting + warehouse_date + local ship-from-warehouse date."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "Для экспорта Excel нужен пакет openpyxl. Установите: pip install openpyxl"
        ) from exc

    details = get_container_modal_details(
        repo,
        user_id=user_id,
        source_id=source_id,
        supply_id=supply_id,
        container_id=container_id,
        container=container,
        client=client,
    )
    warehouse_display = str(details.get("warehouse_date_display") or "").strip()
    if not warehouse_display:
        warehouse_display = str(details.get("warehouse_date") or "").strip() or "—"
    ship_display = str(details.get("moved_to_delivering_at_display") or "").strip()
    if not ship_display:
        ship_display = str(details.get("moved_to_delivering_at") or "").strip() or "—"

    wb = Workbook()
    ws = wb.active
    ws.title = "Состав"
    ws.append(
        [
            "Отправление",
            "Дата склада (Ozon)",
            "Дата отгрузки с нашего склада",
        ]
    )
    for row in details.get("postings") or []:
        if not isinstance(row, dict):
            continue
        pn = str(row.get("posting_number") or "").strip()
        if not pn:
            continue
        ws.append([pn, warehouse_display, ship_display])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 32

    buf = io.BytesIO()
    wb.save(buf)
    payload = buf.getvalue()
    try:
        cid = int(details.get("container_id") or container_id or 0)
    except (TypeError, ValueError):
        cid = int(container_id or 0)
    try:
        num = int(details.get("container_number") or 0)
    except (TypeError, ValueError):
        num = 0
    if num > 0:
        fname = f"GM-{cid}-N{num}-sostav.xlsx"
    else:
        fname = f"GM-{cid}-sostav.xlsx"
    return payload, fname
