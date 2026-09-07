"""Ozon FBS shipment-quality report for support disputes.

Parses Ozon ``general_fbs_rfbs_rating.xlsx`` (posting numbers in column A) and
builds an XLSX with the local date when the posting's supply was moved to
«Доставляются» (ops_log ``move_delivering``), date-only.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import UTC, datetime
from typing import Any

from . import ozon_fbs as oz
from . import ozon_fbs_ops_log as ops_log
from .repository import ReviewRepository

_log = logging.getLogger(__name__)

_HEADER_RE = re.compile(r"^\s*номер\s+отправления\s*$", re.IGNORECASE)
_POSTING_RE = re.compile(r"^\d{6,}-\d{3,}-\d{1,4}$")


def extract_posting_numbers_from_rating_xlsx(content: bytes) -> list[str]:
    """Read column A after the «Номер отправления» header; keep file order."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Для разбора Excel нужен пакет openpyxl. Установите: pip install openpyxl"
        ) from exc
    if not content:
        raise ValueError("Пустой файл")
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise ValueError("Не удалось прочитать Excel-файл") from exc
    try:
        ws = wb.active
        rows = list(ws.iter_rows(min_col=1, max_col=1, values_only=True))
    finally:
        wb.close()

    started = False
    out: list[str] = []
    seen: set[str] = set()
    for row in rows:
        raw = row[0] if row else None
        if raw is None:
            continue
        text = str(raw).strip()
        if not text:
            continue
        if not started:
            if _HEADER_RE.match(text):
                started = True
            continue
        pn = oz.parse_posting_number_query(text) or (
            text if _POSTING_RE.fullmatch(text) else ""
        )
        if not pn:
            continue
        if pn in seen:
            continue
        seen.add(pn)
        out.append(pn)
    if not started:
        raise ValueError(
            "В файле не найден заголовок «Номер отправления» в столбце A"
        )
    if not out:
        raise ValueError("В файле нет номеров отправлений после заголовка")
    return out


def _format_move_date_only(value: object) -> str:
    """MSK calendar date ``DD.MM.YYYY`` from ops_log ``created_at``."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value or "").strip()
        if not text:
            return ""
        if len(text) == 10 and text[4] == "-" and text[7] == "-":
            try:
                return datetime.strptime(text, "%Y-%m-%d").strftime("%d.%m.%Y")
            except ValueError:
                return text
        try:
            dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except Exception:
            return text
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)
    try:
        from zoneinfo import ZoneInfo

        dt = dt.astimezone(ZoneInfo("Europe/Moscow"))
    except Exception:
        dt = dt.astimezone(UTC)
    return dt.strftime("%d.%m.%Y")


def _load_posting_supply_map(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_numbers: list[str],
) -> dict[str, str]:
    """Map posting_number → supply_id for numbers present locally."""
    pns = [str(p).strip() for p in posting_numbers if str(p).strip()]
    if not pns:
        return {}
    oz.ensure_ozon_fbs_tables(repo)
    out: dict[str, str] = {}
    chunk = 500
    with repo._connect() as conn:
        for i in range(0, len(pns), chunk):
            part = pns[i : i + chunk]
            placeholders = ", ".join("?" for _ in part)
            rows = conn.execute(
                repo._sql(
                    f"""
                    SELECT posting_number, supply_id
                    FROM ozon_fbs_postings
                    WHERE user_id = ? AND source_id = ?
                      AND posting_number IN ({placeholders})
                    """
                ),
                (int(user_id), int(source_id), *part),
            ).fetchall()
            for row in rows:
                d = repo._row_to_dict(row)
                pn = str(d.get("posting_number") or "").strip()
                sid = str(d.get("supply_id") or "").strip()
                if pn and sid:
                    out[pn] = sid
    return out


def _load_supply_move_dates(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    supply_ids: list[str],
) -> dict[str, str]:
    """Map supply_id → latest local move-to-delivering date (DD.MM.YYYY)."""
    sids = [str(s).strip() for s in supply_ids if str(s).strip()]
    if not sids:
        return {}
    ops_log.ensure_ozon_fbs_ops_log_table(repo)
    out: dict[str, str] = {}
    chunk = 300
    with repo._connect() as conn:
        for i in range(0, len(sids), chunk):
            part = sids[i : i + chunk]
            placeholders = ", ".join("?" for _ in part)
            rows = conn.execute(
                repo._sql(
                    f"""
                    SELECT supply_id, MAX(created_at) AS moved_at
                    FROM ozon_fbs_ops_log
                    WHERE user_id = ?
                      AND source_id = ?
                      AND action = ?
                      AND supply_id IN ({placeholders})
                    GROUP BY supply_id
                    """
                ),
                (
                    int(user_id),
                    int(source_id),
                    ops_log.ACTION_MOVE_DELIVERING,
                    *part,
                ),
            ).fetchall()
            for row in rows:
                d = repo._row_to_dict(row)
                sid = str(d.get("supply_id") or "").strip()
                if not sid:
                    continue
                out[sid] = _format_move_date_only(d.get("moved_at"))
    return out


def build_support_report_rows(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_numbers: list[str],
) -> list[tuple[str, str]]:
    """Return ``(posting_number, move_date)`` in the same order as input."""
    pns = [str(p).strip() for p in posting_numbers if str(p).strip()]
    supply_by_pn = _load_posting_supply_map(
        repo, user_id=user_id, source_id=source_id, posting_numbers=pns
    )
    supply_ids = sorted({sid for sid in supply_by_pn.values() if sid})
    move_by_supply = _load_supply_move_dates(
        repo, user_id=user_id, source_id=source_id, supply_ids=supply_ids
    )
    rows: list[tuple[str, str]] = []
    for pn in pns:
        sid = supply_by_pn.get(pn) or ""
        date = move_by_supply.get(sid) or ""
        rows.append((pn, date))
    return rows


def build_support_report_xlsx(
    repo: ReviewRepository,
    *,
    user_id: int,
    source_id: int,
    posting_numbers: list[str],
) -> tuple[bytes, str, dict[str, Any]]:
    """Build support XLSX: posting number + local supply move-to-delivering date."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError(
            "Для экспорта Excel нужен пакет openpyxl. Установите: pip install openpyxl"
        ) from exc

    data_rows = build_support_report_rows(
        repo,
        user_id=user_id,
        source_id=source_id,
        posting_numbers=posting_numbers,
    )
    with_date = sum(1 for _pn, d in data_rows if d)
    wb = Workbook()
    ws = wb.active
    ws.title = "Для поддержки"
    ws.append(["Номер отправления", "Дата переноса в доставляются"])
    for pn, date in data_rows:
        ws.append([pn, date])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 32

    buf = io.BytesIO()
    wb.save(buf)
    stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
    fname = f"ozon-fbs-shipment-quality-support-{stamp}.xlsx"
    meta = {
        "total": len(data_rows),
        "with_date": with_date,
        "without_date": len(data_rows) - with_date,
    }
    return buf.getvalue(), fname, meta
